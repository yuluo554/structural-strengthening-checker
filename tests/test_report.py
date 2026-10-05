"""M4 报告导出测试：验算书（FR-D1）+ 核查报告（FR-D2）+ 0 外链守门（FR-D4）。

0 外链口径（plan/02 FR-D4）：产物全部包件关系不得有 external 关系
（不插图片、不加超链接、不引用网络模板——图片内嵌策略下也不允许外链图片）。
"""

import json

import pytest

from structural_strengthening_checker.cli import main
from structural_strengthening_checker.core.clauses import ClauseStore
from structural_strengthening_checker.core.engine import run_calc
from structural_strengthening_checker.core.errors import MissingExtraError
from structural_strengthening_checker.report.docx_export import (
    DISCLAIMER,
    export_calc_book,
    export_check_report,
)

VALID_CFRP_PARAMS = {
    "b": 250.0, "h": 500.0, "fc": 14.3, "As": 1520.0, "fy": 360.0,
    "Afe": 101.0, "Ecf": 240000.0, "M0": 0.0, "M": 180.0,
    "ff": 2300.0, "ef": 0.01, "as": 35.0, "xi_b": 0.518, "Es": 200000.0,
}


def _doc_text(docx_mod, path):
    document = docx_mod.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    return "\n".join(lines), document


def _assert_no_external_links(document):
    """FR-D4：全包件关系无 external（超链接/外链图片/网络模板引用均算外链）。"""
    external = []
    for part in document.part.package.iter_parts():
        rels = getattr(part, "rels", None)
        if not rels:
            continue
        for rel in rels.values():
            if rel.is_external:
                external.append(f"{part.partname} -> {rel.target_ref}")
    assert external == [], f"报告产物存在外链: {external}"


@pytest.fixture
def calc_result(repo_cwd):
    store = ClauseStore("data/knowledge")
    return run_calc("cfrp_beam_flexure", VALID_CFRP_PARAMS, store)


def test_calc_book_export(repo_cwd, tmp_path, calc_result):
    """FR-D1：验算书含结论/输入参数/计算过程/依据条款/签署栏/免责声明，0 外链。"""
    import docx as docx_mod

    out = tmp_path / "验算书.docx"
    export_calc_book(calc_result, out)
    text, document = _doc_text(docx_mod, out)
    assert "验算书：cfrp_beam_flexure" in text
    assert "验算结论：满足" in text
    assert "输入参数" in text and "b" in text  # 参数表回显
    assert "计算过程" in text and "步骤" in text
    assert "依据条款" in text and "GB50367-2013-10.2.3" in text
    assert "签署栏" in text and "计算" in text and "校核" in text and "审核" in text
    assert DISCLAIMER in text
    _assert_no_external_links(document)


def test_check_report_export(repo_cwd, tmp_path):
    """FR-D2：核查报告含概要/明细/证据链/建议/签署栏/免责声明，0 外链。"""
    import docx as docx_mod

    from structural_strengthening_checker.checking.engine import check_plan
    from structural_strengthening_checker.checking.rules import RuleStore
    from structural_strengthening_checker.datagen import gen_plan
    from structural_strengthening_checker.parsing.docx_text import extract_paras
    from structural_strengthening_checker.parsing.param_card_builder import build_param_card

    corpus = tmp_path / "corpus"
    gen_plan.generate(corpus, seed=20261005, count=1)
    plan_path = corpus / "case01_seed20261005.docx"
    paras = extract_paras(plan_path)
    card = build_param_card(paras, doc_name=plan_path.name, doc_sha256="x" * 64)
    report = check_plan(card, RuleStore("data/knowledge"), ClauseStore("data/knowledge"))

    out = tmp_path / "核查报告.docx"
    export_check_report(report, out, card=card)
    text, document = _doc_text(docx_mod, out)
    assert "加固方案核查报告" in text
    assert "核查文档：case01_seed20261005.docx" in text
    assert "sha256" in text
    assert "汇总：不符合 1 / 疑似 0 / 待人工确认 0 / 通过 3" in text
    assert "R-ADH-001" in text
    assert "证据" in text  # 证据链
    assert "签署栏" in text and "核查" in text and "复核" in text and "批准" in text
    assert DISCLAIMER in text
    _assert_no_external_links(document)


def test_report_missing_extra_friendly(monkeypatch, tmp_path, calc_result):
    """缺 report extras：MissingExtraError 带安装提示（CLI 层映射 exit 2，见 test_cli）。"""
    import sys

    from structural_strengthening_checker.report import docx_export

    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(MissingExtraError):
        docx_export.export_calc_book(calc_result, tmp_path / "a.docx")


def test_cli_report_calc_export(repo_cwd, tmp_path, capsys):
    """CLI：ssc report calc → docx 落盘。"""
    params_path = tmp_path / "input.json"
    params_path.write_text(json.dumps(VALID_CFRP_PARAMS, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "book.docx"
    code = main(["report", "calc", "cfrp_beam_flexure", "--input", str(params_path),
                 "--out", str(out)])
    out_text = capsys.readouterr().out
    assert code == 0, out_text
    assert out.is_file()
    assert "验算书 docx 已写入" in out_text


def test_cli_report_check_export(repo_cwd, tmp_path, capsys):
    """CLI：ssc report check（含解析失败降级文档也能出报告）。"""
    import docx as docx_mod

    from structural_strengthening_checker.datagen import gen_plan

    corpus = tmp_path / "corpus"
    gen_plan.generate(corpus, seed=20261005, count=1)
    out = tmp_path / "check.docx"
    code = main(["report", "check", str(corpus / "case01_seed20261005.docx"),
                 "--out", str(out)])
    out_text = capsys.readouterr().out
    assert code == 0, out_text
    assert out.is_file()
    text, document = _doc_text(docx_mod, out)
    assert "R-ADH-001" in text
    _assert_no_external_links(document)

    # 损坏文档：降级报告（全 needs_confirm），不产生 fail、不 traceback
    bad = tmp_path / "损坏方案.docx"
    bad.write_bytes(b"not a docx")
    out2 = tmp_path / "check_bad.docx"
    code = main(["report", "check", str(bad), "--out", str(out2)])
    out_text = capsys.readouterr().out
    assert code == 0, out_text
    text2, document2 = _doc_text(docx_mod, out2)
    assert "待人工确认" in text2 and "不符合 0" in text2
    _assert_no_external_links(document2)


def test_cli_check_report_flag_writes_docx(repo_cwd, tmp_path, capsys):
    """M3 遗留补齐：ssc check --report 目录同时写 JSON + docx。"""
    from structural_strengthening_checker.datagen import gen_plan

    corpus = tmp_path / "corpus"
    gen_plan.generate(corpus, seed=20261005, count=1)
    report_dir = tmp_path / "report"
    code = main(["check", str(corpus / "case01_seed20261005.docx"), "--report", str(report_dir)])
    out_text = capsys.readouterr().out
    assert code == 0, out_text
    assert (report_dir / "check_report.json").is_file()
    assert (report_dir / "check_report.docx").is_file()
    assert "核查报告 docx 已写入" in out_text


def test_cli_report_missing_extra_exit2(repo_cwd, tmp_path, capsys, monkeypatch):
    """CLI 缺 extras → exit 2 友好提示，绝不 traceback。"""
    import sys

    monkeypatch.setitem(sys.modules, "docx", None)
    params_path = tmp_path / "input.json"
    params_path.write_text(json.dumps(VALID_CFRP_PARAMS), encoding="utf-8")
    code = main(["report", "calc", "cfrp_beam_flexure", "--input", str(params_path),
                 "--out", str(tmp_path / "x.docx")])
    out_text = capsys.readouterr().out
    assert code == 2, out_text
    assert "缺少依赖" in out_text
