"""合成方案生成器 v1：固定 seed 确定性 / 注入-规则联动 / 真值语义 / 缺 extras 降级。"""

import json
import sys

import pytest

from structural_strengthening_checker.datagen import gen_plan

pytest.importorskip("docx", reason="dev extras 未安装 python-docx 时跳过生成器测试")


def test_injection_types_map_to_rules():
    """四类注入与规则库 rule_id 一一对应（plan/05 §四联动纪律）。"""
    assert set(gen_plan.INJECTION_RULE_MAP) == set(gen_plan.INJECTION_TYPES)
    assert set(gen_plan.INJECTION_RULE_MAP.values()) == {
        "R-ADH-001",
        "R-ANCH-001",
        "R-APP-001",
        "R-CONST-001",
    }


def test_fixed_seed_reproducible(tmp_path):
    """M1 DoD：固定 seed 两次运行产物一致（docx sha256 + truth 内容）。"""
    truth1 = gen_plan.generate(tmp_path / "run1", seed=42, count=5)
    truth2 = gen_plan.generate(tmp_path / "run2", seed=42, count=5)
    assert [d["doc_sha256"] for d in truth1["docs"]] == [
        d["doc_sha256"] for d in truth2["docs"]
    ]
    t1 = json.loads((tmp_path / "run1" / "truth.json").read_text(encoding="utf-8"))
    t2 = json.loads((tmp_path / "run2" / "truth.json").read_text(encoding="utf-8"))
    assert t1 == t2


def test_baseline_clean_and_count(tmp_path):
    truth = gen_plan.generate(tmp_path, seed=7, count=4)
    docs = truth["docs"]
    assert len(docs) == 5  # 1 基线 + 4 注入
    baseline = docs[0]
    assert baseline["clean_baseline"] is True
    assert baseline["injected"] == []
    assert baseline["expected_non_pass"] == []
    # count=4 → 每类注入各一次，主期望全部 fail、also_expect 为空（v1 互不隐含）
    seen = set()
    for doc in docs[1:]:
        assert doc["clean_baseline"] is False
        assert len(doc["injected"]) == 1
        item = doc["injected"][0]
        assert item["expected_verdict"] == "fail"
        assert item["also_expect"] == []
        assert doc["expected_non_pass"] == [item["rule_id"]]
        seen.add(item["type"])
    assert seen == set(gen_plan.INJECTION_TYPES)


def test_multi_injection_doc_truth_semantics(tmp_path):
    """count>4 时出现多注入组合文档：expected_non_pass = 注入 rule_id 全集。"""
    truth = gen_plan.generate(tmp_path, seed=20261005, count=5)
    combos = [d for d in truth["docs"] if len(d["injected"]) > 1]
    assert combos, "seed=20261005/count=5 应产生多注入组合文档"
    for doc in combos:
        rule_ids = [item["rule_id"] for item in doc["injected"]]
        assert doc["expected_non_pass"] == rule_ids


def test_baseline_doc_contains_compliance_markers(tmp_path):
    """基线文档必须含四条规则的核查对象（鉴定结论/卸载/胶指标/锚固深度）。"""
    gen_plan.generate(tmp_path, seed=1, count=1)
    try:
        import docx as docx_mod
    except ImportError:  # pragma: no cover
        pytest.skip("docx missing")
    document = docx_mod.Document(str(tmp_path / "baseline_seed1.docx"))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "检测鉴定" in text
    assert "卸载" in text
    assert "钢对钢拉伸抗剪强度标准值" in text
    assert "锚固深度" in text
    assert "15.8MPa" in text  # 合规值（≥15）


def test_adhesive_low_value_injected(tmp_path):
    gen_plan.generate(tmp_path, seed=1, count=4)
    try:
        import docx as docx_mod
    except ImportError:  # pragma: no cover
        pytest.skip("docx missing")
    document = docx_mod.Document(str(tmp_path / "case01_seed1.docx"))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "12.0MPa" in text  # 注入的低值（<15 → R-ADH-001 fail）
    assert "15.8MPa" not in text  # 旧行已清除（先清后写，防同参数两处注入打架）


def test_missing_docx_friendly(monkeypatch, tmp_path):
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(Exception) as ei:
        gen_plan.generate(tmp_path, seed=1, count=1)
    assert "[report]" in str(ei.value)
    assert "pip install" in str(ei.value)
